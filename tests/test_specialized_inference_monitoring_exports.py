"""Independent method references and negative-design safeguards."""
from html import escape
import itertools
import zipfile
from xml.etree import ElementTree as ET

import numpy as np
import pandas as pd
import pytest
from scipy import stats
from sklearn.datasets import make_blobs

from khaled_dsml import (
    williams_steiger_overlapping_correlations, effect_size_r_from_z,
    conover_iman_posthoc, polychoric_latent_correlation,
    pelt_mean_shift, ddm_binary_error_monitor, adwin_bounded_mean_monitor,
    save_analysis_xlsx, save_analysis_docx, render_analysis_dashboard,
)


def test_steiger_overlapping_williams_statistic_matches_independent_formula():
    rng=np.random.default_rng(17)
    a=rng.normal(size=(85,3))
    x=a[:,0]
    y=.8*x+.5*a[:,1]
    z=.1*x+.8*a[:,2]
    r=williams_steiger_overlapping_correlations(x,y,z)
    r12,r13,r23=np.corrcoef(np.column_stack([x,y,z]),rowvar=False)[np.triu_indices(3,1)]
    determinant=1-r12*r12-r13*r13-r23*r23+2*r12*r13*r23
    denom=(2*(len(x)-1)*determinant/(len(x)-3)
           + ((r12+r13)**2/4)*(1-r23)**3)
    t=(r12-r13)*np.sqrt((len(x)-1)*(1+r23)/denom)
    assert r["statistic"]==pytest.approx(t,rel=1e-11)
    assert r["p_value"]==pytest.approx(2*stats.t.sf(abs(t),len(x)-3))
    assert r["reject_equal_correlations"]
    assert r["table"].shape==(3,3)


def test_effect_r_formula_signed_and_invalid_overflow():
    r=effect_size_r_from_z(-2.,100)
    assert r["effect_size_r"]==pytest.approx(-.2)
    assert "p_value" not in r
    with pytest.raises(ValueError):effect_size_r_from_z(30,100)
    with pytest.raises(ValueError):effect_size_r_from_z(2,True)


def test_conover_no_ties_rank_formula_and_holm_reference():
    rng=np.random.default_rng(9)
    groups=[rng.normal(loc=i*.7,size=12) for i in range(4)]
    r=conover_iman_posthoc(*groups)
    pooled=np.concatenate(groups)
    ranks=stats.rankdata(pooled)
    n=len(pooled);k=len(groups)
    rank_means=np.array([ranks[i*12:(i+1)*12].mean() for i in range(k)])
    H=stats.kruskal(*groups).statistic
    var=np.var(ranks,ddof=1)*(n-1-H)/(n-k)
    ref=[]
    for i,j in itertools.combinations(range(k),2):
        t=abs(rank_means[i]-rank_means[j])/np.sqrt(var*(1/12+1/12))
        ref.append(2*stats.t.sf(t,n-k))
    np.testing.assert_allclose(r["table"]["raw_p"],ref)
    from statsmodels.stats.multitest import multipletests
    np.testing.assert_allclose(r["table"]["holm_p"],
                               multipletests(ref,method="holm")[1])
    assert len(r["table"])==6
    tied=[np.array([1,2,3,4,5]),np.array([2,6,7,8,9]),
          np.array([10,11,12,13,14])]
    with pytest.raises(ValueError,match="tied"):
        conover_iman_posthoc(*tied)


def test_polychoric_latent_normal_recovers_population_correlation():
    rng=np.random.default_rng(26)
    latent=rng.multivariate_normal([0,0],[[1,.65],[.65,1]],size=950)
    x=np.digitize(latent[:,0],[-.65,.35])
    y=np.digitize(latent[:,1],[-.35,.7])
    r=polychoric_latent_correlation(x,y)
    assert .45<r["polychoric_rho"]<.83
    assert len(r["table"])==3 and r["table"].to_numpy().sum()==len(x)
    assert "p_value" not in r
    with pytest.raises(ValueError):
        polychoric_latent_correlation(x.astype(float)+.5,y)


def _reference_dp(x,penalty,m):
    # Full unpruned dynamic-programming oracle, intentionally independent
    # of the optimized PELT retained-candidate algorithm.
    x=np.asarray(x,float)
    n=len(x)
    f=np.full(n+1,np.inf)
    prev=np.full(n+1,-1,dtype=int)
    f[0]=-penalty
    for t in range(m,n+1):
        for s in range(0,t-m+1):
            if not np.isfinite(f[s]):continue
            segment=x[s:t]
            val=f[s]+((segment-segment.mean())**2).sum()+penalty
            if val<f[t]:
                f[t]=val;prev[t]=s
    splits=[];t=n
    while t>0:
        t=prev[t]
        if t>0:splits.append(t)
    return f[n],sorted(splits)


@pytest.mark.parametrize("seed,penalty,min_size",[
    (2,2.,3),(3,.5,4),(4,8.,5),(8,4.,6),(10,.1,3)
])
def test_pelt_matches_full_dp_on_noisy_and_block_shift_signals(seed,penalty,min_size):
    rng=np.random.default_rng(seed)
    x=np.r_[rng.normal(0,.7,18),rng.normal(3,.7,18)]
    r=pelt_mean_shift(x,penalty=penalty,min_segment_size=min_size)
    score,cps=_reference_dp(x,penalty,min_size)
    assert r["penalized_objective"]==pytest.approx(score,abs=1e-7)
    assert r["change_indices"]==cps
    assert all(item>=min_size for item in np.diff(
        [0,*r["change_indices"],len(x)]))


def test_pelt_constant_series_no_change_and_complex_cases():
    for penalty in (0.1,5.,100.):
        r=pelt_mean_shift(np.ones(60),penalty=penalty,min_segment_size=6)
        assert r["n_changes"]==0
        assert r["penalized_objective"]==pytest.approx(0,abs=1e-6)
    with pytest.raises(ValueError):
        pelt_mean_shift(np.ones(60),penalty=-1)


def test_ddm_zero_error_stream_no_alarm_then_change():
    stream=np.r_[np.zeros(45,dtype=int),np.ones(110,dtype=int)]
    r=ddm_binary_error_monitor(stream,min_instances=30)
    assert r["n_drift_alarms"]>=1
    assert r["drift_indices"][0]>=45
    assert len(r["table"])==len(stream)
    zero=ddm_binary_error_monitor(np.zeros(100,dtype=int))
    assert zero["n_drift_alarms"]==0
    with pytest.raises(ValueError):
        ddm_binary_error_monitor([0,1,2]*20)


def test_adwin_conservative_hoeffding_detects_large_shift():
    x=np.r_[np.zeros(100),np.ones(100)]
    r=adwin_bounded_mean_monitor(x,min_subwindow=12,delta=.05)
    assert r["n_alarms"]>=1
    assert r["alarm_indices"][0]>=100
    assert r["alarm_details"][0]["new_window_start"]>0
    zero=adwin_bounded_mean_monitor(np.zeros(160),min_subwindow=12)
    assert zero["n_alarms"]==0
    with pytest.raises(ValueError):
        adwin_bounded_mean_monitor(np.r_[x[:-1],1.5])


def test_all_new_scientific_functions_have_bilingual_offline_studio():
    rng=np.random.default_rng(29)
    v=rng.normal(size=(110,3))
    rhos=williams_steiger_overlapping_correlations(v[:,0],v[:,1],v[:,2])
    test_results=[
        rhos,
        effect_size_r_from_z(-2.,100),
        conover_iman_posthoc(*[rng.normal(size=10)+j for j in range(3)]),
        pelt_mean_shift(np.r_[np.zeros(40),np.ones(40)*3],penalty=3),
        ddm_binary_error_monitor(np.r_[np.zeros(40,dtype=int),np.ones(40,dtype=int)]),
        adwin_bounded_mean_monitor(np.r_[np.zeros(60),np.ones(60)],min_subwindow=8),
    ]
    for result in test_results:
        html=render_analysis_dashboard(result,theme="amethyst")
        assert 'dir="rtl"' in html
        assert escape(result["interpretation_ar"]) in html
        assert escape(result["interpretation_en"]) in html


def test_excel_export_has_bilingual_text_valid_data_no_executable_formulas(tmp_path):
    from openpyxl import load_workbook
    obj={
        "test":"Formula injection test","n_samples":40,
        "interpretation_ar":"نتيجة اختبار عربي","interpretation_en":"A quoted \"English\" note",
        "table":pd.DataFrame({
            "unsafe":["=2+2","+SUM(A1:A3)","@foo","-cmd"],
            "effect":[.2,.3,.4,.5],
        })
    }
    path=save_analysis_xlsx(obj,tmp_path/"result.xlsx")
    wb=load_workbook(path,data_only=False)
    assert wb.sheetnames==["Summary","Data Table"]
    assert wb["Summary"].sheet_view.rightToLeft
    assert wb["Summary"]["B4"].value
    sh=wb["Data Table"]
    assert sh["A2"].data_type=="s"
    assert sh["A2"].value.startswith("'=2+2")
    assert sh["A3"].value.startswith("'+SUM")
    assert sh["B2"].value==pytest.approx(.2)


def test_docx_export_has_arabic_rtl_and_table_rows(tmp_path):
    obj={"test":"Analysis","n_samples":35,
         "interpretation_ar":"نتيجة باللغة العربية",
         "interpretation_en":"Result in English",
         "table":pd.DataFrame({"group":["A","B"],"mean":[4.2,7.8]})}
    out=save_analysis_docx(obj,tmp_path/"report.docx")
    assert out.exists()
    with zipfile.ZipFile(out) as f:
        assert f.testzip() is None
        payload=f.read("word/document.xml")
        root=ET.fromstring(payload)
    ns={"w":"http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    text=" ".join(t.text or "" for t in root.findall(".//w:t",ns))
    assert "نتيجة باللغة العربية" in text
    assert "Result in English" in text
    assert root.findall(".//w:bidi",ns)
    assert len(root.findall(".//w:tbl",ns))>=2
    with pytest.raises(ValueError):
        save_analysis_docx(obj,tmp_path/"bad.ext")
