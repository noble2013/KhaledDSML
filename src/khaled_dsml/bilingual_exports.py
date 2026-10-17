"""Readable bilingual scientific exports from dictionary-like analysis results.

Save reported metrics, interpretations and summary tables. The helper never
recalculates statistics, adds fictitious significance or embeds remote media.
Optional dependencies are imported only when writing files.
"""
from __future__ import annotations

from collections.abc import Mapping
from pathlib import Path
import math

import pandas as pd


def _validate_result(result, path, expected_suffix):
    if not isinstance(result,Mapping) or not result:
        raise ValueError("result must be a nonempty mapping")
    dest=Path(path)
    if dest.suffix.lower()!=expected_suffix:
        raise ValueError(f"Destination must end with {expected_suffix}")
    table=result.get("table")
    if table is not None and not isinstance(table,pd.DataFrame):
        raise ValueError("result['table'] must be a pandas DataFrame when present")
    if isinstance(table,pd.DataFrame) and (
        table.shape[0]>10000 or table.shape[1]>100
    ):
        raise ValueError("Export table exceeds 10000 rows or 100 columns")
    return dest,table


def _safe_text(value):
    if isinstance(value,(list,tuple,dict,set)):
        return repr(value)[:2000]
    if value is None or pd.isna(value):
        return ""
    return str(value)[:5000]


def _excel_value(value):
    # Refuse to turn arbitrary user text into executable Excel formulas.
    if value is None:
        return None
    if isinstance(value,(bool,int,float)):
        if isinstance(value,float) and not math.isfinite(value):
            return str(value)
        return value
    if pd.isna(value):
        return None
    s=_safe_text(value)
    if s.startswith(("=","+","-","@")):
        return "'"+s
    return s


def save_analysis_xlsx(result, path, *, title="KhaledDSML | تقرير التحليل العلمي"):
    """Export complete bilingual metrics and a bounded table to XLSX.

    Excel formulas are NOT evaluated: text starting with formula trigger
    characters is prefixed with a literal apostrophe. No background network
    requests or user data uploads. Output can contain sensitive user data
    if the caller passes it in result/table.
    """
    dest,table=_validate_result(result,path,".xlsx")
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    except ImportError as exc:
        raise ImportError('Install with: pip install "khaled-dsml[exports]"') from exc
    wb=Workbook()
    ws=wb.active;ws.title="Summary";ws.sheet_view.rightToLeft=True
    ws.merge_cells("A1:D1")
    ws["A1"]=title
    ws["A1"].font=Font(name="Arial",size=16,bold=True,color="FFFFFF")
    ws["A1"].fill=PatternFill("solid",fgColor="262440")
    ws.row_dimensions[1].height=34
    ws["A3"]="المقياس / Metric";ws["B3"]="القيمة / Value"
    scalars=[]
    for key,value in result.items():
        if key=="table" or isinstance(value,(Mapping,pd.DataFrame,list,tuple,set)):
            continue
        scalars.append((str(key),_excel_value(value)))
    for index,(label,val) in enumerate(scalars,start=4):
        ws.cell(index,1,label)
        ws.cell(index,2,val)
        ws.cell(index,1).alignment=Alignment(horizontal="left",vertical="center")
        ws.cell(index,2).alignment=Alignment(horizontal="right" if label.endswith("_ar") else "left",
                                              vertical="top",wrap_text=True)
        if label.startswith("interpretation_"):
            ws.row_dimensions[index].height=55
    ws.column_dimensions["A"].width=37
    ws.column_dimensions["B"].width=87
    ws.freeze_panes="B4"
    for header in ws[3]:
        header.fill=PatternFill("solid",fgColor="2469B6")
        header.font=Font(name="Arial",bold=True,color="FFFFFF")
    if table is not None:
        tab=wb.create_sheet("Data Table")
        tab.sheet_view.rightToLeft=False
        fields=list(map(str,table.columns))
        for col,name in enumerate(fields,1):
            tab.cell(1,col,_excel_value(name))
        for row in range(len(table)):
            for col in range(len(fields)):
                tab.cell(row+2,col+1,_excel_value(table.iat[row,col]))
        tab.freeze_panes="A2"
        tab.auto_filter.ref=f"A1:{tab.cell(1,max(1,len(fields))).column_letter}{len(table)+1}"
        for j in range(1,len(fields)+1):
            tab.column_dimensions[tab.cell(1,j).column_letter].width=24
            hd=tab.cell(1,j)
            hd.fill=PatternFill("solid",fgColor="2469B6")
            hd.font=Font(name="Arial",bold=True,color="FFFFFF")
            hd.alignment=Alignment(horizontal="center",wrap_text=True)
    wb.save(dest)
    return dest


def save_analysis_docx(result,path,*,title="KhaledDSML | تقرير التحليل العلمي"):
    """Export a paginated Arabic RTL / English LTR DOCX research report.

    Preserves reported statistics as provided. DOCX is an editable document,
    not a scientific proof of correctness or a signed research certificate.
    """
    dest,table=_validate_result(result,path,".docx")
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt, RGBColor
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
    except ImportError as exc:
        raise ImportError('Install with: pip install "khaled-dsml[exports]"') from exc
    def paragraph_bidi(par):
        ppr=par._p.get_or_add_pPr()
        direction=OxmlElement("w:bidi")
        direction.set(qn("w:val"),"1")
        ppr.append(direction)
        par.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    def rtl_run(r):
        rpr=r._r.get_or_add_rPr()
        direction=OxmlElement("w:rtl")
        direction.set(qn("w:val"),"1")
        rpr.append(direction)
    doc=Document()
    section=doc.sections[0]
    normal=doc.styles["Normal"]
    normal.font.name="Arial"
    normal.font.size=Pt(11)
    normal.font.color.rgb=RGBColor(35,38,56)
    heading=doc.add_heading(title,level=0)
    paragraph_bidi(heading)
    rtl_run(heading.runs[0])
    ar=str(result.get("interpretation_ar",""))
    en=str(result.get("interpretation_en",""))
    if ar:
        p=doc.add_paragraph()
        paragraph_bidi(p)
        run=p.add_run(ar);rtl_run(run)
    if en:
        doc.add_paragraph(en)
    doc.add_heading("مخرجات التحليل | Analysis metrics",level=1)
    metrics=doc.add_table(rows=1,cols=2)
    metrics.style="Light Shading Accent 1"
    metrics.rows[0].cells[0].text="Metric"
    metrics.rows[0].cells[1].text="Value"
    for key,val in result.items():
        if key=="table" or isinstance(val,(Mapping,pd.DataFrame,list,tuple,set)):
            continue
        row=metrics.add_row().cells
        row[0].text=str(key)
        row[1].text=_safe_text(val)
        if key.endswith("_ar"):
            paragraph_bidi(row[1].paragraphs[0])
            for run in row[1].paragraphs[0].runs:rtl_run(run)
    if table is not None and table.shape[1]:
        doc.add_heading("الجدول العلمي | Evidence table",level=1)
        out=doc.add_table(rows=1,cols=len(table.columns))
        out.style="Light Shading Accent 1"
        for j,k in enumerate(table.columns):
            out.rows[0].cells[j].text=str(k)
        for vals in table.itertuples(index=False,name=None):
            cells=out.add_row().cells
            for j,val in enumerate(vals):
                cells[j].text=_safe_text(val)
    warning=doc.add_paragraph()
    paragraph_bidi(warning)
    rtl_run(warning.add_run("النتائج البحثية تحتاج مراجعة التصميم والافتراضات؛ الدلالة الإحصائية لا تثبت السببية."))
    doc.save(dest)
    return dest
